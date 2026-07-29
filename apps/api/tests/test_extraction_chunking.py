from pathlib import Path

from docx import Document as DocxDocument
from pptx import Presentation

from app.ingestion.chunking import chunk_document
from app.ingestion.extractors import extract_document
from app.ingestion.normalize import normalize_document


def test_extract_markdown_preserves_heading_metadata() -> None:
    extracted = normalize_document(
        extract_document(
            Path("strategy.md"),
            b"# Launch Plan\n\nShip semantic search.\n\n## Risks\n\nHandle privacy carefully.",
            "text/markdown",
        )
    )

    assert extracted.blocks[0].heading == "Launch Plan"
    assert extracted.blocks[1].heading == "Risks"
    assert extracted.blocks[0].source_filename == "strategy.md"


def test_extract_source_code_sets_language_section() -> None:
    extracted = normalize_document(
        extract_document(Path("search.py"), b"def search():\n    pass\n")
    )

    assert extracted.metadata["kind"] == "source_code"
    assert extracted.blocks[0].section == "py"


def test_extract_docx_preserves_headings(tmp_path: Path) -> None:
    path = tmp_path / "notes.docx"
    document = DocxDocument()
    document.add_heading("Meeting Notes", level=1)
    document.add_paragraph("Discuss semantic indexing and retrieval.")
    document.save(path)

    extracted = normalize_document(extract_document(path, path.read_bytes()))

    assert extracted.metadata["kind"] == "docx"
    assert any(block.heading == "Meeting Notes" for block in extracted.blocks)


def test_extract_pptx_preserves_slide_numbers(tmp_path: Path) -> None:
    path = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Search Pipeline"
    presentation.save(path)

    extracted = normalize_document(extract_document(path, path.read_bytes()))

    assert extracted.metadata["kind"] == "pptx"
    assert extracted.blocks[0].slide_number == 1


def test_chunking_is_deterministic_and_token_aware() -> None:
    extracted = normalize_document(
        extract_document(
            Path("memory.txt"),
            ("alpha beta gamma\n\n" * 80).encode("utf-8"),
            "text/plain",
        )
    )

    first = chunk_document(
        extracted,
        model="gemini-embedding-2",
        target_tokens=40,
        overlap_tokens=8,
    )
    second = chunk_document(
        extracted,
        model="gemini-embedding-2",
        target_tokens=40,
        overlap_tokens=8,
    )

    assert [chunk.content_hash for chunk in first] == [chunk.content_hash for chunk in second]
    assert all(chunk.token_count <= 45 for chunk in first)
