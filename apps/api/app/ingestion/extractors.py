from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader

from app.ingestion.models import ExtractedBlock, ExtractedDocument

TEXT_EXTENSIONS = {".txt", ".text", ".csv", ".tsv", ".log"}
MARKDOWN_EXTENSIONS = {".md", ".markdown", ".mdx"}
CODE_EXTENSIONS = {
    ".c",
    ".cpp",
    ".cs",
    ".css",
    ".go",
    ".html",
    ".java",
    ".js",
    ".jsx",
    ".json",
    ".kt",
    ".mjs",
    ".py",
    ".rb",
    ".rs",
    ".sh",
    ".sql",
    ".swift",
    ".toml",
    ".ts",
    ".tsx",
    ".yaml",
    ".yml",
}


class UnsupportedFileTypeError(ValueError):
    pass


def extract_document(path: Path, content: bytes, mime_type: str | None = None) -> ExtractedDocument:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS or (mime_type and mime_type.startswith("text/plain")):
        return _extract_text(path, content, mime_type)
    if suffix in MARKDOWN_EXTENSIONS or mime_type in {"text/markdown", "text/x-markdown"}:
        return _extract_markdown(path, content, mime_type)
    if suffix == ".pdf" or mime_type == "application/pdf":
        return _extract_pdf(path, content, mime_type)
    if (
        suffix == ".docx"
        or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        return _extract_docx(path, content, mime_type)
    if (
        suffix == ".pptx"
        or mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    ):
        return _extract_pptx(path, content, mime_type)
    if suffix in CODE_EXTENSIONS:
        return _extract_code(path, content, mime_type)
    raise UnsupportedFileTypeError(f"Unsupported file type: {path.name}")


def _decode(content: bytes) -> str:
    return content.decode("utf-8", errors="replace")


def _extract_text(path: Path, content: bytes, mime_type: str | None) -> ExtractedDocument:
    return ExtractedDocument(
        blocks=(
            ExtractedBlock(
                text=_decode(content),
                source_filename=path.name,
                block_index=0,
            ),
        ),
        source_filename=path.name,
        mime_type=mime_type,
        metadata={"kind": "plain_text"},
    )


def _extract_markdown(path: Path, content: bytes, mime_type: str | None) -> ExtractedDocument:
    blocks: list[ExtractedBlock] = []
    current_heading: str | None = None
    buffer: list[str] = []
    block_index = 0

    def flush() -> None:
        nonlocal block_index
        text = "\n".join(buffer).strip()
        if text:
            blocks.append(
                ExtractedBlock(
                    text=text,
                    heading=current_heading,
                    section=current_heading,
                    source_filename=path.name,
                    block_index=block_index,
                )
            )
            block_index += 1
        buffer.clear()

    for line in _decode(content).splitlines():
        if line.startswith("#"):
            flush()
            current_heading = line.lstrip("#").strip() or None
        buffer.append(line)
    flush()
    return ExtractedDocument(
        blocks=tuple(blocks),
        source_filename=path.name,
        mime_type=mime_type,
        metadata={"kind": "markdown"},
    )


def _extract_code(path: Path, content: bytes, mime_type: str | None) -> ExtractedDocument:
    return ExtractedDocument(
        blocks=(
            ExtractedBlock(
                text=_decode(content),
                section=path.suffix.lower().lstrip("."),
                source_filename=path.name,
                block_index=0,
            ),
        ),
        source_filename=path.name,
        mime_type=mime_type,
        metadata={"kind": "source_code", "extension": path.suffix.lower()},
    )


def _extract_pdf(path: Path, content: bytes, mime_type: str | None) -> ExtractedDocument:
    reader = PdfReader(io_bytes(content))
    blocks = []
    for index, page in enumerate(reader.pages):
        blocks.append(
            ExtractedBlock(
                text=page.extract_text() or "",
                page_number=index + 1,
                source_filename=path.name,
                block_index=index,
            )
        )
    return ExtractedDocument(
        blocks=tuple(blocks),
        source_filename=path.name,
        mime_type=mime_type,
        metadata={"kind": "pdf", "page_count": len(reader.pages)},
    )


def _extract_docx(path: Path, content: bytes, mime_type: str | None) -> ExtractedDocument:
    document = DocxDocument(io_bytes(content))
    blocks: list[ExtractedBlock] = []
    current_heading: str | None = None
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.lower().startswith("heading"):
            current_heading = text
        blocks.append(
            ExtractedBlock(
                text=text,
                heading=current_heading,
                section=current_heading,
                source_filename=path.name,
                block_index=index,
            )
        )
    return ExtractedDocument(
        blocks=tuple(blocks),
        source_filename=path.name,
        mime_type=mime_type,
        metadata={"kind": "docx", "paragraph_count": len(document.paragraphs)},
    )


def _extract_pptx(path: Path, content: bytes, mime_type: str | None) -> ExtractedDocument:
    presentation = Presentation(io_bytes(content))
    blocks: list[ExtractedBlock] = []
    for slide_index, slide in enumerate(presentation.slides):
        lines: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = str(shape.text).strip()
                if text:
                    lines.append(text)
        blocks.append(
            ExtractedBlock(
                text="\n".join(lines),
                slide_number=slide_index + 1,
                source_filename=path.name,
                block_index=slide_index,
            )
        )
    return ExtractedDocument(
        blocks=tuple(blocks),
        source_filename=path.name,
        mime_type=mime_type,
        metadata={"kind": "pptx", "slide_count": len(presentation.slides)},
    )


def io_bytes(content: bytes) -> BytesIO:
    return BytesIO(content)
